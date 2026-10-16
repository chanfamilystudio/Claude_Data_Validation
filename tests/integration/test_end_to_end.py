"""Synthetic end-to-end suite: the whole workflow through the CLI only, against a fresh database.

migrate -> validate -> run -> export -> edit workbook -> import -> sign off -> run again -> revoke.
The staging data comes from tests/fixtures/synthetic.sql (every break type planted) or
synthetic_clean.sql (nothing to find). Employees are snapshot rules, payroll is periodic; the
periodic rows are dated September 2026, hence ``--period 2026-09``.
"""

from __future__ import annotations

import uuid
from pathlib import Path

import pyodbc
import pytest
import yaml
from openpyxl import load_workbook
from typer.testing import CliRunner

from recon.cli import app
from recon.db import Database

from .conftest import run_sql_file
from .test_executor import _reset_recon_logger  # noqa: F401  (fixture)

pytestmark = pytest.mark.integration

FIXTURES = Path(__file__).resolve().parent.parent / "fixtures"
PERIOD = "2026-09"
BREAK_TYPES = {"missing_left", "missing_right", "value_mismatch", "mapping_missing", "mapping_mismatch",
               "quality_violation", "total_mismatch"}
DOMAINS = ("employees", "payroll")

# (rule_id, break_type, business_key, field_name) of every break synthetic.sql plants, and nothing else.
PLANTED = {
    ("emp_master_match", "missing_right", "000006", None),
    ("emp_master_match", "missing_left", "000007", None),
    ("emp_master_match", "missing_left", "ABC123", None),
    ("emp_master_match", "value_mismatch", "000003", "annual_salary"),
    ("emp_master_match", "value_mismatch", "000004", "national_id"),  # PII, shown as ***
    ("emp_master_match", "value_mismatch", "000005", "cost_center"),  # effective-dated mapping row change
    ("emp_master_match", "mapping_missing", "000008", "cost_center"),
    ("emp_master_match", "quality_violation", "000010", "__key__"),  # duplicate key
    ("emp_cost_center_map", "mapping_missing", "CC9", "cost_center"),
    ("emp_cost_center_map", "mapping_mismatch", "CC3", "cost_center"),
    ("gp_emp_quality", "quality_violation", "ABC123", "not_null:DEPRTMNT"),
    ("gp_emp_quality", "quality_violation", "ABC123", "like:EMPLOYID"),
    ("gp_emp_quality", "quality_violation", "000010", "unique:EMPLOYID"),
    ("gp_emp_quality", "quality_violation", "000011", "accepted_values:INACTIVE"),
    ("gp_emp_quality", "quality_violation", "000011", "future_hire_dates"),  # custom_sql
    ("gp_emp_quality", "quality_violation", "000009", "references:DEPRTMNT"),
    # control-total groups are keyed by the GP account index (101 = A6100, 102 = A6200)
    ("pay_jrnl_totals", "total_mismatch", "101", "net_amount"),
    ("pay_jrnl_totals", "total_mismatch", "102", "net_amount"),
    ("pay_jrnl_totals", "total_mismatch", "102", "line_count"),
    ("pay_jrnl_totals", "total_mismatch", "UNMAPPED:A9999", "net_amount"),
    ("pay_jrnl_totals", "total_mismatch", "UNMAPPED:A9999", "line_count"),
}


@pytest.fixture
def e2e_db(server_dsn):
    """A fresh, empty database (no ``recon`` schema yet): ``recon db migrate`` has to create it."""
    name = f"recon_e2e_{uuid.uuid4().hex[:8]}"
    master = pyodbc.connect(server_dsn, autocommit=True)
    try:
        master.execute(f"CREATE DATABASE [{name}]")
        database = Database(f"{server_dsn.rstrip(';')};Database={name}")
        try:
            yield database
        finally:
            database.close()
    finally:
        try:
            master.execute(f"ALTER DATABASE [{name}] SET SINGLE_USER WITH ROLLBACK IMMEDIATE")
            master.execute(f"DROP DATABASE [{name}]")
        finally:
            master.close()


def _load(db: Database, script: str) -> None:
    run_sql_file(db, FIXTURES / "synthetic_teardown.sql")  # a teardown must work on an empty database too
    run_sql_file(db, FIXTURES / script)


@pytest.fixture
def synthetic(e2e_db):
    _load(e2e_db, "synthetic.sql")
    yield e2e_db
    run_sql_file(e2e_db, FIXTURES / "synthetic_teardown.sql")
    assert e2e_db.scalar("SELECT COUNT(*) FROM sys.schemas WHERE name IN ('stg_wd', 'stg_gp', 'xref')") == 0


@pytest.fixture
def project(tmp_fixture, e2e_db, monkeypatch):
    """A tmp copy of the fixture project that signs as the connection's own login."""
    monkeypatch.setenv("RECON_DB", e2e_db.conn_str)
    me = e2e_db.current_user()
    (tmp_fixture / "config" / "signers.yaml").write_text(
        yaml.safe_dump({"signers": {d: [me] for d in DOMAINS}}), encoding="utf-8")
    return tmp_fixture


def _cli(runner: CliRunner, project: Path, *args: str):
    return runner.invoke(app, [*args, "--root", str(project)])


def _edit_workbook(path: Path) -> tuple[int, int]:
    """Accept the first break (with a comment) and resolve every other one (with comments)."""
    wb = load_workbook(path)
    ws = wb["Breaks"]
    header = [c.value for c in ws[1]]
    status, comment = header.index("Status") + 1, header.index("New comment") + 1
    ids = [row[0].value for row in ws.iter_rows(min_row=2) if row[0].value is not None]
    assert len(ids) >= len(PLANTED)
    for n, row in enumerate(ws.iter_rows(min_row=2), start=1):
        if row[0].value is None:
            continue
        accepted = n == 1
        ws.cell(row=row[0].row, column=status).value = "accepted" if accepted else "resolved"
        ws.cell(row=row[0].row, column=comment).value = (
            "known difference, approved by finance" if accepted else "corrected at source")
    wb.save(path)
    return len(ids), 1


def test_end_to_end_workflow(synthetic, project, _reset_recon_logger):  # noqa: F811
    db = synthetic
    runner = CliRunner()

    # migrate (fresh database) and validate
    res = _cli(runner, project, "db", "migrate")
    assert res.exit_code == 0 and res.output.startswith("Applied: "), res.output
    assert _cli(runner, project, "db", "migrate").output.strip() == "Up to date"
    res = _cli(runner, project, "rules", "validate")
    assert res.exit_code == 0 and res.output.strip() == "OK: 4 rules", res.output

    # run everything: breaks, no errors
    res = _cli(runner, project, "run", "--all", "--period", PERIOD)
    assert res.exit_code == 1, res.output
    for line in ("emp_cost_center_map: fail", "emp_master_match: fail", "gp_emp_quality: fail",
                 "pay_jrnl_totals: fail"):
        assert any(out.startswith(line) for out in res.output.splitlines()), res.output
    breaks = db.query("SELECT rule_id, break_type, business_key, field_name, left_value, right_value, status"
                      " FROM recon.[break]")
    assert {b["break_type"] for b in breaks} == BREAK_TYPES
    assert {b["status"] for b in breaks} == {"open"}
    # custom_sql keys come straight from the rule's SQL, so CHAR padding is not trimmed
    assert {(b["rule_id"], b["break_type"], b["business_key"].strip(), b["field_name"]) for b in breaks} == PLANTED
    pii = [b for b in breaks if b["rule_id"] == "emp_master_match" and b["field_name"] == "national_id"]
    assert [(b["business_key"], b["left_value"], b["right_value"]) for b in pii] == [("000004", "***", "***")]
    assert db.scalar("SELECT COUNT(*) FROM recon.run WHERE status = 'succeeded'") == 1

    # nothing can be signed off while breaks are open
    res = _cli(runner, project, "signoff", "--domain", "payroll", "--period", PERIOD, "--comment", "too early")
    assert res.exit_code == 2 and "breaks are open or investigating" in res.output, res.output

    # export, edit, import
    res = _cli(runner, project, "export", "--all", "--period", PERIOD)
    assert res.exit_code == 0, res.output
    book = Path(res.output.strip().splitlines()[-1])
    assert book.is_file() and book.parent == project / "exports"
    total, accepted = _edit_workbook(book)
    res = _cli(runner, project, "import", str(book))
    assert res.exit_code == 0 and res.output.strip() == f"applied {total}, rejected 0", res.output
    by_status = {r["status"]: r["n"] for r in db.query(
        "SELECT status, COUNT(*) AS n FROM recon.[break] GROUP BY status")}
    assert by_status == {"accepted": accepted, "resolved": total - accepted}
    assert db.scalar("SELECT COUNT(*) FROM recon.break_comment") == total

    # sign off every domain
    evidence = []
    for domain in DOMAINS:
        res = _cli(runner, project, "signoff", "--domain", domain, "--period", PERIOD,
                   "--comment", f"{domain} September reviewed")
        assert res.exit_code == 0, res.output
        path = Path(res.output.strip().rsplit(": ", 1)[1])
        assert path.is_file() and path.parent == project / "evidence" / domain / PERIOD
        evidence.append(path)
    assert db.scalar("SELECT COUNT(*) FROM recon.signoff WHERE status = 'signed'") == len(DOMAINS)

    # a later run warns that the periods are signed off, and still reports the breaks
    res = _cli(runner, project, "run", "--all", "--period", PERIOD)
    assert res.exit_code == 1, res.output
    for domain in DOMAINS:
        assert f"warning: {domain} {PERIOD} is signed off; new breaks are flagged post_signoff" in res.output

    # revoke
    for domain in DOMAINS:
        res = _cli(runner, project, "signoff", "revoke", "--domain", domain, "--period", PERIOD,
                   "--reason", "late adjustment")
        assert res.exit_code == 0 and res.output.strip() == f"Revoked {domain} {PERIOD}", res.output
    assert db.scalar("SELECT COUNT(*) FROM recon.signoff WHERE status = 'revoked'") == len(DOMAINS)
    assert all(p.is_file() for p in evidence)  # frozen evidence is kept


def test_clean_data_has_no_breaks(e2e_db, project, _reset_recon_logger):  # noqa: F811
    _load(e2e_db, "synthetic_clean.sql")
    runner = CliRunner()
    assert _cli(runner, project, "db", "migrate").exit_code == 0
    res = _cli(runner, project, "run", "--all", "--period", PERIOD)
    assert res.exit_code == 0, res.output
    assert e2e_db.scalar("SELECT COUNT(*) FROM recon.[break]") == 0
    results = e2e_db.query("SELECT rule_id, result FROM recon.run_check ORDER BY rule_id")
    assert {r["rule_id"]: r["result"] for r in results} == {
        "emp_cost_center_map": "pass", "emp_master_match": "pass", "gp_emp_quality": "pass",
        "pay_jrnl_totals": "pass",
    }
    # the run actually compared rows
    assert e2e_db.scalar("SELECT rows_left FROM recon.run_check WHERE rule_id = 'emp_master_match'") == 3
