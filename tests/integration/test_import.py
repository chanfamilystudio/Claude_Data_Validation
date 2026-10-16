"""``recon import`` end to end against SQL Server: run -> export -> edit -> import."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from recon.cli import app

from .test_executor import _reset_recon_logger, staging  # noqa: F401  (fixtures)

pytestmark = pytest.mark.integration


def _edit_first_row(path, **cells):
    wb = load_workbook(path)
    ws = wb["Breaks"]
    header = [c.value for c in ws[1]]
    for name, value in cells.items():
        ws.cell(row=2, column=header.index(name) + 1).value = value
    break_id = ws.cell(row=2, column=1).value
    wb.save(path)
    return break_id


def test_cli_import_round_trip(staging, dsn, tmp_fixture, monkeypatch, _reset_recon_logger):  # noqa: F811
    db = staging
    monkeypatch.setenv("RECON_DB", dsn)
    db.execute("UPDATE stg_gp.UPR00100 SET SALARY = 1 WHERE EMPLOYID = '000002'")
    runner = CliRunner()
    res = runner.invoke(app, ["run", "--domain", "employees", "--period", "2026-09", "--root", str(tmp_fixture)])
    assert res.exit_code == 1, res.output
    book = next((tmp_fixture / "exports").glob("recon_employees_2026-09_*.xlsx"))
    break_id = _edit_first_row(book, Status="Investigating", Owner="alice", **{"New comment": "checking payroll"})

    res = runner.invoke(app, ["import", str(book), "--dry-run", "--root", str(tmp_fixture)])
    assert res.exit_code == 0 and res.output.strip() == "dry run: applied 1, rejected 0", res.output
    assert db.scalar("SELECT COUNT(*) FROM recon.import_batch") == 0

    res = runner.invoke(app, ["import", str(book), "--root", str(tmp_fixture)])
    assert res.exit_code == 0 and res.output.strip() == "applied 1, rejected 0", res.output
    row = db.query("SELECT status, owner, version FROM recon.[break] WHERE break_id = ?", break_id)[0]
    assert (row["status"], row["owner"], row["version"]) == ("investigating", "alice", 2)
    batch = db.query("SELECT * FROM recon.import_batch")[0]
    assert (batch["file_name"], batch["rows_applied"], batch["rows_rejected"]) == (book.name, 1, 0)
    assert batch["imported_by"] == db.current_user()
    comment = db.query("SELECT [text], author, import_batch_id FROM recon.break_comment WHERE break_id = ?",
                       break_id)[0]
    assert (comment["text"], comment["author"], comment["import_batch_id"]) == (
        "checking payroll", db.current_user(), batch["import_batch_id"])
    audit = db.query("SELECT field, old_value, new_value, source, import_batch_id FROM recon.audit_log"
                     " WHERE source = 'import' ORDER BY field")
    assert [(a["field"], a["old_value"], a["new_value"], a["source"]) for a in audit] == [
        ("owner", None, "alice", "import"), ("status", "open", "investigating", "import")]
    assert {a["import_batch_id"] for a in audit} == {batch["import_batch_id"]}

    res = runner.invoke(app, ["import", str(book), "--root", str(tmp_fixture)])
    assert res.exit_code == 2 and "file already imported" in res.output
    assert db.scalar("SELECT COUNT(*) FROM recon.import_batch") == 1


def test_cli_import_rejects_exit_1(staging, dsn, tmp_fixture, monkeypatch, _reset_recon_logger):  # noqa: F811
    db = staging
    monkeypatch.setenv("RECON_DB", dsn)
    db.execute("UPDATE stg_gp.UPR00100 SET SALARY = 1 WHERE EMPLOYID = '000002'")
    runner = CliRunner()
    assert runner.invoke(app, ["run", "--domain", "employees", "--period", "2026-09",
                               "--root", str(tmp_fixture)]).exit_code == 1
    book = next((tmp_fixture / "exports").glob("recon_employees_2026-09_*.xlsx"))
    break_id = _edit_first_row(book, Status="resolved")  # no comment
    res = runner.invoke(app, ["import", str(book), "--root", str(tmp_fixture)])
    rejects = book.with_name(book.stem + "_rejects.xlsx")
    assert res.exit_code == 1, res.output
    assert res.output.splitlines() == ["applied 0, rejected 1", f"rejects: {rejects}"]
    ws = load_workbook(rejects)["Rejects"]
    assert [c.value for c in ws[2]][-1] == "comment required for resolved"
    assert db.scalar("SELECT status FROM recon.[break] WHERE break_id = ?", break_id) == "open"
    assert db.scalar("SELECT COUNT(*) FROM recon.import_batch") == 0
