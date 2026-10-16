"""Excel import: export with Task 11, edit with openpyxl, import (InMemoryRepository as the database)."""

from __future__ import annotations

import hashlib

import pytest
from openpyxl import Workbook, load_workbook
from typer.testing import CliRunner

from fakes import InMemoryRepository
from recon import cli
from recon.cli import app
from recon.excel import layout
from recon.excel.export import export_workbook
from recon.excel.importer import ImportResult, WorkbookError, import_workbook
from recon.model import UserEdit
from repository_contract import add_rule, change, create_break, observed, open_run

P = "2026-09"
COL = {name: i for i, name in enumerate(layout.BREAK_COLUMNS, start=1)}


@pytest.fixture
def repo():
    r = InMemoryRepository(user="CORP\\importer")
    add_rule(r)
    return r


@pytest.fixture
def breaks(repo):
    run = open_run(repo)
    return [create_break(repo, run, bk="E1"), create_break(repo, run, bk="E2"), create_break(repo, run, bk="E3")]


@pytest.fixture
def exported(repo, breaks, tmp_path):
    return export_workbook(repo, domain="employees", period_label=P, run_id=None, out_path=tmp_path / "book.xlsx")


def _edit(path, edits, *, mutate=None):
    """Set cells of the Breaks rows by break_id: ``edits = {break_id: {"Status": ..., ...}}``."""
    wb = load_workbook(path)
    ws = wb["Breaks"]
    for row in range(2, ws.max_row + 1):
        values = edits.get(ws.cell(row=row, column=1).value)
        for name, value in (values or {}).items():
            ws.cell(row=row, column=COL[name]).value = value
    if mutate:
        mutate(wb)
    wb.save(path)


def _rejects(path):
    ws = load_workbook(path.with_name(path.stem + "_rejects.xlsx"))["Rejects"]
    header = [c.value for c in ws[1]]
    return [dict(zip(header, [c.value for c in row], strict=True)) for row in ws.iter_rows(min_row=2)]


def _audit(repo):
    return [a for a in repo.audit if a["source"] == "import"]


def test_valid_edit_applied(repo, breaks, exported):
    b = breaks[0]
    _edit(exported, {b.break_id: {"Status": "investigating", "Owner": "alice"}})
    res = import_workbook(repo, exported)
    assert (res.applied, res.rejected, res.rejects_path) == (1, 0, None) and res.batch_id is not None
    got = repo.get_break(b.break_id)
    assert (got.status, got.owner, got.version) == ("investigating", "alice", b.version + 1)
    assert sorted((a["field"], a["new_value"]) for a in _audit(repo)) == [
        ("owner", "alice"), ("status", "investigating")]
    assert repo.batches[0]["file_name"] == "book.xlsx"
    assert repo.batches[0]["file_sha256"] == hashlib.sha256(exported.read_bytes()).hexdigest()
    assert not exported.with_name("book_rejects.xlsx").exists()
    # untouched breaks are not edits
    assert repo.get_break(breaks[1].break_id).version == breaks[1].version


def test_comment_only_edit(repo, breaks, exported):
    b = breaks[0]
    _edit(exported, {b.break_id: {"New comment": "  looking into it "}})
    res = import_workbook(repo, exported)
    assert res.applied == 1
    assert [c["text"] for c in repo.comments] == ["looking into it"]
    assert repo.get_break(b.break_id).status == "open"


def test_comment_required(repo, breaks, exported):
    b = breaks[0]
    _edit(exported, {b.break_id: {"Status": "resolved"}})
    res = import_workbook(repo, exported)
    assert (res.applied, res.rejected, res.batch_id) == (0, 1, None)
    assert [r["Reason"] for r in _rejects(exported)] == ["comment required for resolved"]
    assert repo.get_break(b.break_id).status == "open" and repo.batches == []


def test_resolved_with_comment_applies(repo, breaks, exported):
    b = breaks[0]
    _edit(exported, {b.break_id: {"Status": "resolved", "New comment": "fixed in GP"}})
    assert import_workbook(repo, exported).applied == 1
    assert repo.get_break(b.break_id).status == "resolved"


def test_changed_since_export(repo, breaks, exported):
    b = breaks[0]
    repo.apply_user_edits([UserEdit(b.break_id, b.version, "investigating", None, None)], "x.xlsx", "f" * 64, 0)
    _edit(exported, {b.break_id: {"Status": "accepted", "New comment": "ok"}})
    res = import_workbook(repo, exported)
    assert (res.applied, res.rejected) == (0, 1)
    assert [r["Reason"] for r in _rejects(exported)] == ["changed since export"]
    assert repo.get_break(b.break_id).status == "investigating"


def test_nightly_reseen_does_not_block(repo, breaks, exported):
    b = breaks[0]
    run2 = open_run(repo)
    repo.apply_engine_changes([change("seen", obs=observed("E1"), existing=b)], run2)
    assert repo.get_break(b.break_id).version == b.version  # a re-see leaves the counter alone
    _edit(exported, {b.break_id: {"Status": "investigating"}})
    res = import_workbook(repo, exported)
    assert (res.applied, res.rejected) == (1, 0)
    assert repo.get_break(b.break_id).status == "investigating"


def test_rows_sorted_deleted_and_blank(repo, breaks, exported):
    b1, b2, b3 = breaks

    def rearrange(wb):
        ws = wb["Breaks"]
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
        ws.delete_rows(2, ws.max_row)
        rows.reverse()
        rows = [r for r in rows if r[0] != b2.break_id]  # a deleted row
        rows.insert(1, [None] * len(layout.BREAK_COLUMNS))  # a blank row
        for r in rows:
            ws.append(r)

    _edit(exported, {b1.break_id: {"Status": "investigating"}, b3.break_id: {"Owner": "bob"}}, mutate=rearrange)
    res = import_workbook(repo, exported)
    assert (res.applied, res.rejected) == (2, 0)
    assert repo.get_break(b1.break_id).status == "investigating"
    assert repo.get_break(b3.break_id).owner == "bob"
    assert repo.get_break(b2.break_id).version == b2.version


def test_same_file_twice(repo, breaks, exported):
    _edit(exported, {breaks[0].break_id: {"Status": "investigating"}})
    import_workbook(repo, exported)
    with pytest.raises(WorkbookError, match="^file already imported$"):
        import_workbook(repo, exported)
    with pytest.raises(WorkbookError, match="file already imported"):
        import_workbook(repo, exported, dry_run=True)


def test_dry_run_applies_nothing(repo, breaks, exported):
    _edit(exported, {breaks[0].break_id: {"Status": "investigating"}, breaks[1].break_id: {"Status": "resolved"}})
    res = import_workbook(repo, exported, dry_run=True)
    assert (res.applied, res.rejected, res.batch_id) == (1, 1, None)
    assert res.rejects_path == exported.with_name("book_rejects.xlsx") and res.rejects_path.is_file()
    assert repo.batches == [] and repo.get_break(breaks[0].break_id).status == "open"
    # nothing was recorded, so the real import still works
    assert import_workbook(repo, exported).applied == 1


def test_rejects_file_contents(repo, breaks, exported):
    b1, b2, b3 = breaks
    _edit(exported, {
        b1.break_id: {"Status": "resolved"},
        b2.break_id: {"Status": "bogus", "Owner": "x"},
        b3.break_id: {"Owner": "o" * 129},
    })
    old = exported.with_name("book_rejects.xlsx")
    old.write_text("stale")
    res = import_workbook(repo, exported)
    assert (res.applied, res.rejected, res.rejects_path) == (0, 3, old)
    ws = load_workbook(old)["Rejects"]
    assert ws.title == "Rejects" and [c.value for c in ws[1]] == [*layout.BREAK_COLUMNS, "Reason"]
    rows = {r["break_id"]: r for r in _rejects(exported)}
    assert rows[b1.break_id]["Reason"] == "comment required for resolved"
    assert rows[b2.break_id]["Reason"] == "invalid status 'bogus'"
    assert rows[b3.break_id]["Reason"] == "owner longer than 128 characters"
    assert rows[b1.break_id]["Status"] == "resolved" and rows[b2.break_id]["business_key"] == "E2"
    assert rows[b2.break_id]["Owner"] == "x" and rows[b3.break_id]["Owner"] == "o" * 129


def test_partial_apply_with_rejects(repo, breaks, exported):
    _edit(exported, {breaks[0].break_id: {"Status": "investigating"}, breaks[1].break_id: {"Status": "resolved"}})
    res = import_workbook(repo, exported)
    assert (res.applied, res.rejected) == (1, 1)
    assert repo.batches[0]["rows_applied"] == 1 and repo.batches[0]["rows_rejected"] == 1


def test_only_rejects_writes_no_batch(repo, breaks, exported):
    _edit(exported, {breaks[0].break_id: {"Status": "resolved"}})
    res = import_workbook(repo, exported)
    assert res.batch_id is None and repo.batches == [] and res.rejects_path.is_file()


def test_status_case_insensitive_and_trimmed(repo, breaks, exported):
    _edit(exported, {breaks[0].break_id: {"Status": "  Investigating "}})
    assert import_workbook(repo, exported).applied == 1
    assert repo.get_break(breaks[0].break_id).status == "investigating"


def test_invalid_status_shows_trimmed_original(repo, breaks, exported):
    _edit(exported, {breaks[0].break_id: {"Status": "  Done "}})
    import_workbook(repo, exported)
    assert [r["Reason"] for r in _rejects(exported)] == ["invalid status 'Done'"]


def test_blank_status_and_owner_mean_no_change(repo, breaks, exported):
    b = breaks[0]
    repo.apply_user_edits([UserEdit(b.break_id, b.version, "investigating", "alice", None)], "x.xlsx", "f" * 64, 0)
    path = export_workbook(repo, domain="employees", period_label=P, run_id=None, out_path=exported.parent / "b2.xlsx")
    _edit(path, {b.break_id: {"Status": None, "Owner": "  ", "New comment": "note"}})
    res = import_workbook(repo, path)
    assert (res.applied, res.rejected) == (1, 0)
    got = repo.get_break(b.break_id)
    assert (got.status, got.owner) == ("investigating", "alice")
    assert len(_audit(repo)) == 2  # only the setup edit: the comment-only import adds no status / owner audit


def test_unchanged_values_are_not_edits(repo, breaks, exported):
    _edit(exported, {breaks[0].break_id: {"Status": "OPEN"}})  # same status, different case
    res = import_workbook(repo, exported)
    assert (res.applied, res.rejected, res.batch_id) == (0, 0, None)
    assert repo.batches == []


def test_duplicate_break_id(repo, breaks, exported):
    b = breaks[0]

    def duplicate(wb):
        ws = wb["Breaks"]
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value == b.break_id:
                values = [c.value for c in ws[row]]
                break
        ws.append(values)

    _edit(exported, {b.break_id: {"Status": "investigating"}}, mutate=duplicate)
    res = import_workbook(repo, exported)
    assert (res.applied, res.rejected) == (1, 1)
    assert [r["Reason"] for r in _rejects(exported)] == ["duplicate break_id"]
    assert repo.get_break(b.break_id).status == "investigating"


@pytest.mark.parametrize("make_id", [lambda i: float(i), lambda i: str(i), lambda i: f" {i} ", lambda i: f"{i}.0"])
def test_break_id_numeric_forms(repo, breaks, exported, make_id):
    b = breaks[0]

    def retype(wb):
        ws = wb["Breaks"]
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value == b.break_id:
                ws.cell(row=row, column=1).value = make_id(b.break_id)

    _edit(exported, {b.break_id: {"Status": "investigating"}}, mutate=retype)
    assert import_workbook(repo, exported).applied == 1


@pytest.mark.parametrize("bad", ["abc", 1.5, 99999, True])
def test_unknown_break_id(repo, breaks, exported, bad):
    def retype(wb):
        ws = wb["Breaks"]
        ws.cell(row=2, column=1).value = bad
        ws.cell(row=2, column=COL["Status"]).value = "investigating"

    _edit(exported, {}, mutate=retype)
    res = import_workbook(repo, exported)
    assert (res.applied, res.rejected) == (0, 1)
    assert [r["Reason"] for r in _rejects(exported)] == ["unknown break_id"]


def test_break_id_missing_from_meta_is_unknown(repo, breaks, exported):
    run2 = open_run(repo)
    extra = create_break(repo, run2, bk="E9")  # in the database but not in this file's _meta

    def add_row(wb):
        wb["Breaks"].append([extra.break_id, *[None] * (len(layout.BREAK_COLUMNS) - 1)])
        wb["Breaks"].cell(row=wb["Breaks"].max_row, column=COL["Status"]).value = "investigating"

    _edit(exported, {}, mutate=add_row)
    import_workbook(repo, exported)
    assert [r["Reason"] for r in _rejects(exported)] == ["unknown break_id"]
    assert repo.get_break(extra.break_id).status == "open"


def test_check_order_status_before_version(repo, breaks, exported):
    b = breaks[0]
    repo.apply_user_edits([UserEdit(b.break_id, b.version, "investigating", None, None)], "x.xlsx", "f" * 64, 0)
    _edit(exported, {b.break_id: {"Status": "nope"}})
    import_workbook(repo, exported)
    assert [r["Reason"] for r in _rejects(exported)] == ["invalid status 'nope'"]


def test_transition_not_allowed(repo, breaks, exported):
    b = breaks[0]
    _edit(exported, {b.break_id: {"Status": "accepted", "New comment": "known"}})
    import_workbook(repo, exported)
    p2 = export_workbook(repo, domain="employees", period_label=P, run_id=None, out_path=exported.parent / "b2.xlsx")
    _edit(p2, {b.break_id: {"Status": "resolved", "New comment": "done"}})
    res = import_workbook(repo, p2)
    reason = _rejects(p2)[0]["Reason"]
    assert res.rejected == 1 and reason == "transition accepted->resolved not allowed"


def test_not_a_recon_workbook(tmp_path, repo):
    plain = tmp_path / "plain.xlsx"
    Workbook().save(plain)
    with pytest.raises(WorkbookError, match="^not a recon export workbook$"):
        import_workbook(repo, plain)
    garbage = tmp_path / "garbage.xlsx"
    garbage.write_text("not a zip")
    with pytest.raises(WorkbookError, match="^not a recon export workbook$"):
        import_workbook(repo, garbage)


def test_wrong_header_is_not_a_recon_workbook(repo, breaks, exported):
    def rename(wb):
        wb["Breaks"].cell(row=1, column=COL["Owner"]).value = "Assignee"

    _edit(exported, {}, mutate=rename)
    with pytest.raises(WorkbookError, match="^not a recon export workbook$"):
        import_workbook(repo, exported)


def test_missing_meta_sheet(repo, breaks, exported):
    _edit(exported, {}, mutate=lambda wb: wb.remove(wb["_meta"]))
    with pytest.raises(WorkbookError, match="^not a recon export workbook$"):
        import_workbook(repo, exported)


def test_unsupported_schema_version(repo, breaks, exported):
    def bump(wb):
        ws = wb["_meta"]
        for row in range(1, 6):
            if ws.cell(row=row, column=1).value == "schema_version":
                ws.cell(row=row, column=2).value = "2"

    _edit(exported, {breaks[0].break_id: {"Status": "investigating"}}, mutate=bump)
    with pytest.raises(WorkbookError, match="^unsupported schema version 2$"):
        import_workbook(repo, exported)
    assert repo.batches == []


def test_apply_error_propagates(repo, breaks, exported, monkeypatch):
    _edit(exported, {breaks[0].break_id: {"Status": "investigating"}})

    def boom(*a, **k):
        raise RuntimeError("db down")

    monkeypatch.setattr(repo, "apply_user_edits", boom)
    with pytest.raises(RuntimeError, match="db down"):
        import_workbook(repo, exported)


def _change_between_validation_and_apply(repo, monkeypatch, break_id):
    """Someone else edits ``break_id`` after the importer validated it, before the batch is applied."""
    original = repo.apply_user_edits

    def racing(edits, **kw):
        current = repo.get_break(break_id)
        original([UserEdit(break_id, current.version, None, "racer", None)], file_name="other.xlsx",
                 file_sha256="9" * 64, rows_rejected=0)
        return original(edits, **kw)

    monkeypatch.setattr(repo, "apply_user_edits", racing)


def test_version_conflict_at_apply_rolls_back_whole_batch(repo, breaks, exported, monkeypatch):
    a, b = breaks[0], breaks[1]
    _edit(exported, {a.break_id: {"Status": "investigating", "New comment": "x"},
                     b.break_id: {"Owner": "alice"}})
    _change_between_validation_and_apply(repo, monkeypatch, b.break_id)
    with pytest.raises(WorkbookError, match=rf"^changed since export \(break {b.break_id}\); re-run import$"):
        import_workbook(repo, exported)
    assert repo.get_break(a.break_id).status == "open"  # nothing of the batch applied
    assert repo.comments == []
    assert [x["file_name"] for x in repo.batches] == ["other.xlsx"]
    assert not repo.import_seen(hashlib.sha256(exported.read_bytes()).hexdigest())


def test_import_result_is_a_value():
    assert ImportResult(1, 0, None, 3).batch_id == 3


# -- CLI -----------------------------------------------------------------------------------------
@pytest.fixture
def wired(repo, monkeypatch, fixture_root):
    class FakeDb:
        def close(self):
            pass

    monkeypatch.setenv("RECON_DB", "Driver={x};Server=none")
    monkeypatch.setattr(cli, "Database", lambda conn: FakeDb())
    monkeypatch.setattr(cli, "SqlRepository", lambda db: repo)
    return fixture_root


def _invoke(path, root, *extra):
    return CliRunner().invoke(app, ["import", str(path), "--root", str(root), *extra])


def test_cli_applied_exit_0(wired, repo, breaks, exported):
    _edit(exported, {breaks[0].break_id: {"Status": "investigating"}})
    res = _invoke(exported, wired)
    assert res.exit_code == 0 and res.output.strip() == "applied 1, rejected 0"


def test_cli_rejects_exit_1(wired, repo, breaks, exported):
    _edit(exported, {breaks[0].break_id: {"Status": "investigating"}, breaks[1].break_id: {"Status": "resolved"}})
    res = _invoke(exported, wired)
    assert res.exit_code == 1
    assert res.output.splitlines() == ["applied 1, rejected 1", f"rejects: {exported.with_name('book_rejects.xlsx')}"]


def test_cli_dry_run(wired, repo, breaks, exported):
    _edit(exported, {breaks[0].break_id: {"Status": "investigating"}, breaks[1].break_id: {"Status": "resolved"}})
    res = _invoke(exported, wired, "--dry-run")
    assert res.exit_code == 1 and res.output.splitlines()[0] == "dry run: applied 1, rejected 1"
    assert repo.batches == []


def test_cli_workbook_error_exit_2(wired, repo, breaks, exported):
    _edit(exported, {breaks[0].break_id: {"Status": "investigating"}})
    assert _invoke(exported, wired).exit_code == 0
    res = _invoke(exported, wired)
    assert res.exit_code == 2 and "file already imported" in res.output


def test_cli_version_conflict_exit_2(wired, repo, breaks, exported, monkeypatch):
    b = breaks[0]
    _edit(exported, {b.break_id: {"Status": "investigating"}})
    _change_between_validation_and_apply(repo, monkeypatch, b.break_id)
    res = _invoke(exported, wired)
    assert res.exit_code == 2
    assert res.output.strip() == f"changed since export (break {b.break_id}); re-run import"


def test_cli_missing_file_exit_2(wired, tmp_path):
    res = _invoke(tmp_path / "nope.xlsx", wired)
    assert res.exit_code == 2 and "nope.xlsx" in res.output


def test_cli_unexpected_error_exit_2(wired, repo, breaks, exported, monkeypatch):
    def boom(*a, **k):
        raise ValueError("kaput")

    monkeypatch.setattr(cli, "import_workbook", boom)
    res = _invoke(exported, wired)
    assert res.exit_code == 2 and "kaput" in res.output


def test_cli_config_error_exit_2(exported, tmp_path, monkeypatch):
    monkeypatch.delenv("RECON_DB", raising=False)
    res = _invoke(exported, tmp_path)
    assert res.exit_code == 2 and "connection.yaml" in res.output
