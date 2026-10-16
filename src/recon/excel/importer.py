"""Import an edited export workbook (spec 7.2): validate every edited row, apply the valid ones."""

from __future__ import annotations

import hashlib
import io
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from recon.breaks import validate_user_transition
from recon.excel import layout
from recon.model import UserEdit
from recon.repository import Repository

REJECTS_SHEET = "Rejects"
REASON = "Reason"
OWNER_MAX = 128

_NOT_RECON = "not a recon export workbook"


class WorkbookError(Exception):
    """The file as a whole cannot be imported."""


@dataclass(frozen=True)
class ImportResult:
    applied: int
    rejected: int
    rejects_path: Path | None
    batch_id: int | None


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value: Any) -> str | None:
    """Trimmed text of a cell, ``None`` when blank."""
    if _blank(value):
        return None
    return str(value).strip()


def _to_id(value: Any) -> int | None:
    """Integer break_id from an int, an integral float or a numeric string; ``None`` otherwise."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    if isinstance(value, str):
        try:
            return _to_id(int(value.strip()))
        except ValueError:
            try:
                return _to_id(float(value.strip()))
            except ValueError:
                return None
    return None


def _open(data: bytes) -> Workbook:
    try:
        wb = load_workbook(io.BytesIO(data))
    except Exception as exc:  # not a zip, not OOXML, corrupt parts ...
        raise WorkbookError(_NOT_RECON) from exc
    if layout.META not in wb.sheetnames or layout.BREAKS not in wb.sheetnames:
        raise WorkbookError(_NOT_RECON)
    header = [c.value for c in next(wb[layout.BREAKS].iter_rows(min_row=1, max_row=1), ())]
    while header and header[-1] is None:
        header.pop()
    if header != layout.BREAK_COLUMNS:
        raise WorkbookError(_NOT_RECON)
    return wb


def _meta(ws: Worksheet) -> tuple[str | None, dict[int, int]]:
    """Schema version and ``{break_id: version at export}`` from the ``_meta`` sheet."""
    keys = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
            for r in range(1, layout.META_HEADER_ROW)}
    raw = keys.get("schema_version")
    versions: dict[int, int] = {}
    for break_id, version in ws.iter_rows(min_row=layout.META_HEADER_ROW + 1, max_col=2, values_only=True):
        bid, ver = _to_id(break_id), _to_id(version)
        if bid is not None and ver is not None:
            versions[bid] = ver
    return (None if raw is None else str(raw).strip()), versions


def _write_rejects(path: Path, rows: list[tuple[list[Any], str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = REJECTS_SHEET
    ws.append([*layout.BREAK_COLUMNS, REASON])
    for values, reason in rows:
        ws.append([*values, reason])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # never let user text turn into a formula
    tmp = path.with_name(path.name + ".tmp")
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except BaseException:
        tmp.unlink(missing_ok=True)
        raise


def import_workbook(repo: Repository, path: Path, *, dry_run: bool = False) -> ImportResult:
    """Validate every edited row of ``path`` and apply the valid ones in one repository call.

    Rows are matched by ``break_id`` (never by position).  A blank Status / Owner means "no change".
    With ``dry_run`` every check runs and the rejects file is written, but nothing is applied.
    """
    path = Path(path)
    data = path.read_bytes()
    wb = _open(data)
    version, exported_versions = _meta(wb[layout.META])
    if version != layout.SCHEMA_VERSION:
        raise WorkbookError(f"unsupported schema version {version}")
    digest = hashlib.sha256(data).hexdigest()
    if repo.import_seen(digest):
        raise WorkbookError("file already imported")

    status_i, owner_i, comment_i = (layout.BREAK_COLUMNS.index(c) for c in layout.EDITABLE)
    width = len(layout.BREAK_COLUMNS)
    edits: list[UserEdit] = []
    rejects: list[tuple[list[Any], str]] = []
    seen: set[int] = set()

    for row in wb[layout.BREAKS].iter_rows(min_row=2, max_col=width, values_only=True):
        values = list(row) + [None] * (width - len(row))
        if _blank(values[0]):
            continue
        break_id = _to_id(values[0])
        current = repo.get_break(break_id) if break_id in exported_versions else None
        if break_id is None or current is None:
            rejects.append((values, "unknown break_id"))
            continue

        raw_status, owner, comment = _text(values[status_i]), _text(values[owner_i]), _text(values[comment_i])
        status = raw_status.lower() if raw_status is not None else None
        new_status = status if status is not None and status != current.status else None
        new_owner = owner if owner is not None and owner != current.owner else None
        duplicate = break_id in seen
        seen.add(break_id)
        if new_status is None and new_owner is None and comment is None:
            continue  # not an edit
        if duplicate:
            rejects.append((values, "duplicate break_id"))
            continue

        reason = None
        if new_status is not None and new_status not in layout.STATUSES:
            reason = f"invalid status '{raw_status}'"
        elif exported_versions[break_id] != current.version:
            reason = "changed since export"
        else:
            reason = validate_user_transition(current.status, new_status or current.status, comment)
            if reason is None and new_owner is not None and len(new_owner) > OWNER_MAX:
                reason = f"owner longer than {OWNER_MAX} characters"
        if reason is not None:
            rejects.append((values, reason))
            continue
        edits.append(UserEdit(break_id, exported_versions[break_id], new_status, new_owner, comment))

    rejects_path = None
    if rejects:
        rejects_path = path.with_name(f"{path.stem}_rejects.xlsx")
        _write_rejects(rejects_path, rejects)
    batch_id = None
    if edits and not dry_run:
        batch_id = repo.apply_user_edits(edits, file_name=path.name, file_sha256=digest,
                                         rows_rejected=len(rejects))
    return ImportResult(len(edits), len(rejects), rejects_path, batch_id)
