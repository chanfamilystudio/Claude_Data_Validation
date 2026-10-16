"""Write the run / evidence workbook (spec 7.1) from a ``Repository``."""

from __future__ import annotations

import os
import uuid
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from recon.excel import layout
from recon.model import BreakRecord, Comment, RunRow
from recon.repository import Repository

_UNLOCKED = Protection(locked=False)
_HEADER_FONT = Font(bold=True)
_BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
_EDIT_FILL = PatternFill("solid", fgColor="FFF2CC")
_COMMENT_LIMIT = 5
CELL_MAX = 32_767  # Excel's limit on characters in one cell
SQL_PART = 32_000  # generated SQL is split into parts of this many characters
_WIDTHS = {"break_id": 10, "business_key": 22, "left": 22, "right": 22, "comment_history": 50,
           "Status": 14, "Owner": 18, "New comment": 40}


def export_path(export_dir: Path, domain: str | None, period_label: str, run_id: int | None) -> Path:
    """``<export_dir>/recon_<domain|all>_<period>_<run_id or "latest">.xlsx``."""
    suffix = run_id if run_id is not None else "latest"
    return Path(export_dir) / f"recon_{domain or 'all'}_{period_label}_{suffix}.xlsx"


def _clean(value: Any) -> Any:
    """Cell value for ``value``: numbers and booleans as such, everything else as literal text."""
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, int | float):
        return value
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    return _fit(ILLEGAL_CHARACTERS_RE.sub("", str(value)))


def _fit(text: str) -> str:
    """``text`` cut to one Excel cell, ending in a visible `` …[truncated N chars]`` marker when cut."""
    if len(text) <= CELL_MAX:
        return text
    dropped = len(text) - CELL_MAX
    while True:  # the marker's own length depends on the number in it
        marker = f" …[truncated {dropped} chars]"
        keep = CELL_MAX - len(marker)
        if len(text) - keep == dropped:
            return text[:keep] + marker
        dropped = len(text) - keep


def _sql_parts(sql: str | None) -> list[str | None]:
    """``sql`` in chunks of ``SQL_PART`` characters (concatenated in part order they give ``sql`` back)."""
    if not sql:
        return [sql]
    return [sql[i:i + SQL_PART] for i in range(0, len(sql), SQL_PART)]


def _put(ws: Worksheet, row: int, col: int, value: Any):
    """Write one cell; strings are always stored as text, never evaluated as formulas."""
    cell = ws.cell(row=row, column=col)
    cell.value = _clean(value)
    if isinstance(cell.value, str):
        cell.data_type = "s"
    return cell


def _write_header(ws: Worksheet, columns: list[str], row: int = 1) -> None:
    for i, name in enumerate(columns, start=1):
        cell = _put(ws, row, i, name)
        cell.font = _HEADER_FONT
        ws.column_dimensions[cell.column_letter].width = max(12, _WIDTHS.get(name, len(name) + 4))


def _write_rows(ws: Worksheet, rows: list[list[Any]], first_row: int = 2) -> None:
    for r, values in enumerate(rows, start=first_row):
        for c, value in enumerate(values, start=1):
            _put(ws, r, c, value)


def _protect(ws: Worksheet) -> None:
    ws.protection.sheet = True  # no password: a convenience lock, enforcement happens on import
    ws.protection.autoFilter = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def _comment_history(comments: list[Comment]) -> str | None:
    latest = sorted(comments, key=lambda c: c.created_at)[-_COMMENT_LIMIT:]
    if not latest:
        return None
    return "\n".join(f"{c.author} {c.created_at:%Y-%m-%d}: {c.text}" for c in latest)


def _sort_key(b: BreakRecord) -> tuple:
    return layout.SEVERITY_ORDER.get(b.severity, len(layout.SEVERITY_ORDER)), b.break_type, b.business_key, \
        b.break_id


def _summary_sheet(ws: Worksheet, repo: Repository, run: RunRow | None, run_id: int | None,
                   domain: str | None, period_label: str) -> None:
    if run is not None:
        meta = [run.run_id, run.domain or "all", run.period_label, run.status, run.run_by, run.started_at,
                run.finished_at]
    elif run_id is not None:
        meta = [run_id, domain or "all", period_label, None, None, None, None]
    else:
        meta = [layout.NO_RUN, domain or "all", period_label, None, None, None, None]
    for r, (key, value) in enumerate(zip(layout.SUMMARY_RUN_KEYS, meta, strict=True), start=1):
        _put(ws, r, 1, key).font = _HEADER_FONT
        _put(ws, r, 2, value)
    header_row = len(layout.SUMMARY_RUN_KEYS) + 2
    _write_header(ws, layout.SUMMARY_CHECK_COLUMNS, header_row)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["G"].width = 60
    if run_id is not None:
        checks = repo.checks_for_run(run_id, domain)
        _write_rows(ws, [[c.rule_id, c.severity, c.result, c.rows_left, c.rows_right, c.break_count,
                          c.error_message]
                         for c in checks], header_row + 1)


def _totals_sheet(ws: Worksheet, repo: Repository, run_id: int | None, domain: str | None) -> None:
    _write_header(ws, layout.TOTALS_COLUMNS)
    if run_id is None:
        return
    for r, t in enumerate(repo.totals_for_run(run_id, domain), start=2):
        _write_rows(ws, [[t.rule_id, t.group_key, t.measure, t.left_value, t.right_value, t.diff, t.tolerance_applied,
                          t.within_tolerance]], r)
        if not t.within_tolerance:
            for c in range(1, len(layout.TOTALS_COLUMNS) + 1):
                ws.cell(row=r, column=c).fill = _BAD_FILL


def _breaks_sheet(ws: Worksheet, breaks: list[BreakRecord], history: dict[int, list[Comment]],
                  evidence: bool) -> None:
    _write_header(ws, layout.BREAK_COLUMNS)
    _write_rows(ws, [[
        b.break_id, b.rule_id, b.severity, b.domain, b.period_label, b.business_key, b.break_type,
        b.field_name, b.left_value, b.right_value, b.diff, b.first_seen_run_id, b.last_seen_run_id,
        b.post_signoff, _comment_history(history.get(b.break_id, [])), b.status, b.owner, None,
    ] for b in breaks])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(layout.BREAK_COLUMNS)).column_letter}{len(breaks) + 1}"
    history_col = layout.BREAK_COLUMNS.index("comment_history") + 1
    for r in range(2, len(breaks) + 2):
        ws.cell(row=r, column=history_col).alignment = Alignment(wrap_text=True, vertical="top")
    if evidence or not breaks:
        return
    first, last = 2, len(breaks) + 1
    for name in layout.EDITABLE:
        col = layout.BREAK_COLUMNS.index(name) + 1
        for r in range(first, last + 1):
            cell = ws.cell(row=r, column=col)
            cell.protection = _UNLOCKED
            cell.fill = _EDIT_FILL
    status = ws.cell(row=1, column=layout.BREAK_COLUMNS.index("Status") + 1).column_letter
    validation = DataValidation(type="list", formula1='"' + ",".join(layout.STATUSES) + '"', allow_blank=False,
                                showErrorMessage=True, errorTitle="Status",
                                error="Pick one of: " + ", ".join(layout.STATUSES))
    validation.add(f"{status}{first}:{status}{last}")
    ws.add_data_validation(validation)


def _meta_sheet(ws: Worksheet, run_id: int | None, exported_by: str, breaks: list[BreakRecord]) -> None:
    values = [str(uuid.uuid4()), run_id, layout.SCHEMA_VERSION, exported_by,
              datetime.now(UTC).isoformat(timespec="seconds")]
    for r, (key, value) in enumerate(zip(layout.META_KEYS, values, strict=True), start=1):
        _put(ws, r, 1, key)
        _put(ws, r, 2, value)
    _write_header(ws, layout.META_HEADER, layout.META_HEADER_ROW)
    _write_rows(ws, [[b.break_id, b.version] for b in breaks], layout.META_HEADER_ROW + 1)
    ws.sheet_state = "veryHidden"


def export_workbook(repo: Repository, *, domain: str | None, period_label: str, run_id: int | None,
                    out_path: Path, evidence: bool = False) -> Path:
    """Write the workbook for ``domain`` (``None`` = all domains) and ``period_label`` to ``out_path``.

    ``run_id`` selects the run shown on Summary / Totals / SQL and the break set (breaks resolved in
    that run are included); when ``None`` the latest run covering ``domain`` is shown and the breaks
    resolved in *that* run are included.  Evidence mode lists every break, adds the
    Comments / Audit / SQL sheets and leaves nothing editable.
    """
    if evidence and domain is None:
        raise ValueError("evidence export requires a domain")
    out_path = Path(out_path)
    exported_by = repo.current_user()

    run = repo.get_run(run_id) if run_id is not None else repo.latest_run(domain, period_label)
    shown_run = run_id if run_id is not None else (run.run_id if run is not None else None)

    if evidence:
        breaks = repo.breaks_for(domain, period_label)
    else:
        breaks = repo.breaks_for(domain, period_label, statuses={"open", "investigating", "accepted"},
                                 resolved_in_run=shown_run)
    breaks = sorted(breaks, key=_sort_key)
    comments = repo.comments_for([b.break_id for b in breaks])

    wb = Workbook()
    summary = wb.active
    summary.title = layout.SUMMARY
    _summary_sheet(summary, repo, run, shown_run, domain, period_label)
    _totals_sheet(wb.create_sheet(layout.TOTALS), repo, shown_run, domain)
    _breaks_sheet(wb.create_sheet(layout.BREAKS), breaks, comments, evidence)
    if evidence:
        ws = wb.create_sheet(layout.COMMENTS)
        _write_header(ws, layout.COMMENT_COLUMNS)
        _write_rows(ws, [[c.break_id, c.author, c.created_at, c.text]
                         for b in breaks for c in sorted(comments.get(b.break_id, []), key=lambda c: c.created_at)])
        ws.column_dimensions["D"].width = 80
        ws = wb.create_sheet(layout.AUDIT)
        _write_header(ws, layout.AUDIT_COLUMNS)
        _write_rows(ws, [[a.entity, a.entity_id, a.field, a.old_value, a.new_value, a.changed_by, a.changed_at,
                          a.source] for a in repo.audit_for(domain, period_label)])
        ws = wb.create_sheet(layout.SQL)
        _write_header(ws, layout.SQL_COLUMNS)
        ws.column_dimensions["C"].width = 120
        if shown_run is not None:
            _write_rows(ws, [[c.rule_id, n, part] for c in repo.checks_for_run(shown_run, domain)
                             for n, part in enumerate(_sql_parts(c.generated_sql), start=1)])
    _meta_sheet(wb.create_sheet(layout.META), shown_run, exported_by, breaks)
    for ws in wb:
        _protect(ws)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = out_path.with_name(out_path.name + ".tmp")
    try:
        wb.save(tmp)
        os.replace(tmp, out_path)
    except BaseException:
        tmp.unlink(missing_ok=True)
        raise
    return out_path
